from pathlib import Path

from openpyxl import load_workbook

from contingency_solver.core import CandidateLine, ConductorModel
from contingency_solver.mock import simulate_results
from contingency_solver import storage
from contingency_solver.storage import create_working_case_copy, export_excel


def test_excel_export_basic(tmp_path: Path) -> None:
    results = simulate_results([CandidateLine(1, "A", 2, "B", 115.0, "115", 5.0)])
    conductors = {"115": ConductorModel("115", "test", 115, 1, 0.1, 0.2, "susceptance_per_mile", 0.0, 100, 110, 120)}
    path = tmp_path / "results.xlsx"
    export_excel(results, path, conductors, {"mock_mode": True})
    workbook = load_workbook(path)
    assert "Candidate Results" in workbook.sheetnames
    assert "Run Summary" in workbook.sheetnames


def test_create_working_case_copy(monkeypatch, tmp_path: Path) -> None:
    source = tmp_path / "Example Case.pwb"
    source.write_text("case-data", encoding="utf-8")
    working_dir = tmp_path / "working"
    monkeypatch.setattr(storage, "USER_DATA", tmp_path / "user_data")
    monkeypatch.setattr(storage, "LOG_DIR", tmp_path / "user_data" / "logs")
    monkeypatch.setattr(storage, "EXPORT_DIR", tmp_path / "user_data" / "exports")
    monkeypatch.setattr(storage, "WORKING_CASE_DIR", working_dir)

    copy_path = create_working_case_copy(source)

    assert copy_path != source
    assert copy_path.parent == working_dir
    assert copy_path.suffix == ".pwb"
    assert copy_path.read_text(encoding="utf-8") == "case-data"
