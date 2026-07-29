from pathlib import Path

from openpyxl import load_workbook

from contingency_solver.services.export_service import ExportService
from contingency_solver.services.mock_data import simulate_results
from contingency_solver.models.candidate import CandidateLine
from contingency_solver.models.conductor import ConductorModel


def test_excel_export_basic(tmp_path: Path) -> None:
    results = simulate_results([CandidateLine(1, "A", 2, "B", 115.0, "115", 5.0)])
    conductors = {"115": ConductorModel("115", "test", 115, 1, 0.1, 0.2, "susceptance_per_mile", 0.0, 100, 110, 120)}
    path = tmp_path / "results.xlsx"
    ExportService().export_results_excel(results, path, conductors, {"mock_mode": True})
    workbook = load_workbook(path)
    assert "Candidate Results" in workbook.sheetnames
    assert "Run Summary" in workbook.sheetnames
